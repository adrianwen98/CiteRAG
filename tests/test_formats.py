from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient

import api_main
from file_utils import DocumentNoTextError, read_document_pages
from langchain_file_utils import load_and_split_documents
from api_main import app


client = TestClient(app)


def make_pdf(path: Path, pages: list[str | None]) -> Path:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def test_markdown_and_text_pdf_are_split_with_page_metadata(tmp_path):
    markdown = tmp_path / "guide.md"
    markdown.write_text("# 标题\n\n这是一段 Markdown 文档。", encoding="utf-8")
    md_docs = load_and_split_documents(str(markdown), 300, 50)
    assert md_docs and md_docs[0].metadata["file_type"] == ".md"
    assert "page" not in md_docs[0].metadata

    pdf = make_pdf(tmp_path / "manual.pdf", ["第一页内容", "第二页内容", None])
    pages = read_document_pages(str(pdf))
    assert [page for page, _ in pages] == [1, 2]
    pdf_docs = load_and_split_documents(str(pdf), 300, 50)
    assert {doc.metadata["page"] for doc in pdf_docs} == {1, 2}
    assert all(doc.metadata["file_type"] == ".pdf" for doc in pdf_docs)


def test_word_paragraphs_and_tables_are_extracted(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "manual.docx"
    document = DocxDocument()
    document.add_paragraph("第一段：产品说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "价格"
    table.cell(1, 0).text = "A1"
    table.cell(1, 1).text = "99"
    document.save(path)

    docs = load_and_split_documents(str(path), 300, 50)
    assert docs and docs[0].metadata["file_type"] == ".docx"
    assert "page" not in docs[0].metadata
    text = "\n".join(doc.page_content for doc in docs)
    assert "产品说明" in text and "型号" in text and "99" in text


def test_scanned_pdf_without_text_is_rejected(tmp_path):
    scanned = make_pdf(tmp_path / "scanned.pdf", [None])
    with pytest.raises(DocumentNoTextError):
        read_document_pages(str(scanned))


def test_upload_accepts_markdown_and_rejects_unknown_extension(tmp_path, monkeypatch):
    monkeypatch.setattr(api_main, "DATA_DIR", tmp_path)
    monkeypatch.setattr(api_main, "index_file_langchain", lambda file_path: "文件索引成功，并写入1个片段")

    response = client.post(
        "/documents/upload",
        files={"file": ("notes.md", "# 标题\n\n内容".encode("utf-8"), "text/markdown")},
    )
    assert response.status_code == 200
    assert (tmp_path / "notes.md").is_file()

    rejected = client.post(
        "/documents/upload",
        files={"file": ("legacy.doc", b"not supported", "application/msword")},
    )
    assert rejected.status_code == 400


def test_scanned_pdf_upload_is_rejected_and_rolled_back(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setattr(api_main, "DATA_DIR", data_dir)
    scanned = make_pdf(tmp_path / "source-scanned.pdf", [None])
    response = client.post(
        "/documents/upload",
        files={"file": ("scanned-upload.pdf", scanned.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 400
    assert not (data_dir / "scanned-upload.pdf").exists()


def test_pdf_citation_includes_page_number():
    from langchain_core.documents import Document
    from langchain_rag_service import format_documents

    text = format_documents([
        Document(page_content="page text", metadata={"source": "/tmp/manual.pdf", "file_type": ".pdf", "page": 2}),
        Document(page_content="text", metadata={"source": "/tmp/notes.md", "file_type": ".md"}),
    ])
    assert "manual.pdf 第2页" in text
    assert "notes.md" in text and "notes.md 第" not in text


def test_index_accepts_markdown_under_data_dir(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setattr(api_main, "DATA_DIR", data_dir)
    monkeypatch.setattr(api_main, "index_file_langchain", lambda file_path: "文件索引成功，并写入1个片段")
    note = data_dir / "indexed.md"
    note.write_text("# 标题\n\n内容", encoding="utf-8")

    response = client.post("/documents/index", json={"file_path": str(note)})
    assert response.status_code == 200
    assert "1个片段" in response.json()["message"]

    rejected = client.post("/documents/index", json={"file_path": str(data_dir / "bad.docx")})
    assert rejected.status_code in {400, 404}


def test_root_serves_single_page_ui():
    response = client.get("/")
    assert response.status_code == 200
    assert "text/html" in response.headers["content-type"]
    assert "RAG 知识库问答" in response.text
